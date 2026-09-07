import csv
import io
import re
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from app.schemas.imports import ImportParticipant, ImportPreview

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 5000
MAX_PDF_PAGES = 10

_EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.IGNORECASE)
_PHONE_RE = re.compile(r"(?<!\w)(?:\+?\d[\d\s()./\-]{5,}\d)(?!\w)")
_LEADING_INDEX_RE = re.compile(r"^\s*(?:[-•·]|\d{1,4}[.)\-:]?)\s*")
_HEADER_WORDS = {
    "name", "names", "participant", "participants", "person", "people",
    "first name", "last name", "firstname", "lastname", "email", "e-mail", "mail",
    "phone", "telephone", "mobile", "contact",
    "vorname", "nachname", "name", "teilnehmer", "teilnehmerin", "teilnehmerinnen",
    "e-mail", "telefon", "handy", "kontakt",
    "nom", "prenom", "prénom", "courriel", "telephone", "téléphone",
    "nombre", "apellido", "correo", "telefono", "teléfono",
    "nome", "cognome", "telefono", "telefon",
}


class ImportParseError(ValueError):
    pass


def _clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).replace("\u00a0", " ").split()).strip()


def _looks_like_phone(value: str) -> bool:
    digits = sum(char.isdigit() for char in value)
    if digits < 6 or digits > 20:
        return False
    return bool(_PHONE_RE.fullmatch(value.strip()))


def _looks_like_header(values: list[str]) -> bool:
    normalized = {value.casefold().strip(" :") for value in values if value}
    return bool(normalized) and len(normalized & _HEADER_WORDS) >= min(2, len(normalized))


def _participant_from_cells(cells: list[object]) -> ImportParticipant | None:
    values = [_clean(value) for value in cells]
    values = [value for value in values if value]
    if not values or _looks_like_header(values):
        return None

    email: str | None = None
    phone: str | None = None
    name_parts: list[str] = []

    for value in values:
        email_match = _EMAIL_RE.search(value)
        if email_match and email is None:
            email = email_match.group(0).strip().lower()
            remainder = _EMAIL_RE.sub(" ", value).strip(" ,;|-/")
            if remainder and not _looks_like_phone(remainder):
                name_parts.append(remainder)
            continue
        if _looks_like_phone(value) and phone is None:
            phone = value
            continue
        if value.isdigit() and len(value) <= 4:
            continue
        name_parts.append(value)

    name = " ".join(name_parts)
    name = _LEADING_INDEX_RE.sub("", name).strip(" ,;|-/")
    name = " ".join(name.split())
    if not name or len(name) > 200:
        return None
    if sum(char.isalpha() for char in name) < 2:
        return None

    try:
        return ImportParticipant(name=name, email=email, phone=phone)
    except ValueError:
        return None


def _parse_lines(text: str) -> list[ImportParticipant]:
    participants: list[ImportParticipant] = []
    for raw_line in text.splitlines():
        line = " ".join(raw_line.replace("\u00a0", " ").split()).strip()
        if not line or len(line) > 500:
            continue
        if line.casefold().startswith(("page ", "seite ", "pagina ", "página ")):
            continue

        email_match = _EMAIL_RE.search(line)
        email = email_match.group(0).strip().lower() if email_match else None
        without_email = _EMAIL_RE.sub(" ", line)

        phone: str | None = None
        for candidate in _PHONE_RE.findall(without_email):
            if _looks_like_phone(candidate):
                phone = " ".join(candidate.split())
                without_email = without_email.replace(candidate, " ", 1)
                break

        name = _LEADING_INDEX_RE.sub("", without_email)
        name = re.sub(r"\s*[|;,]\s*", " ", name)
        name = " ".join(name.split()).strip(" ,;|-/")
        if not name or len(name) > 200 or sum(char.isalpha() for char in name) < 2:
            continue
        if name.casefold().strip(" :") in _HEADER_WORDS:
            continue

        try:
            participants.append(ImportParticipant(name=name, email=email, phone=phone))
        except ValueError:
            continue
        if len(participants) >= MAX_IMPORT_ROWS:
            break
    return participants


def _deduplicate(participants: list[ImportParticipant]) -> list[ImportParticipant]:
    result: list[ImportParticipant] = []
    seen: set[tuple[str, str, str]] = set()
    for participant in participants:
        key = (
            participant.name.casefold(),
            (participant.email or "").casefold(),
            re.sub(r"\D", "", participant.phone or ""),
        )
        if key in seen:
            continue
        seen.add(key)
        result.append(participant)
    return result


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ImportParseError("Text encoding not supported")


def _parse_csv(data: bytes) -> list[ImportParticipant]:
    text = _decode_text(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    participants: list[ImportParticipant] = []
    for row in reader:
        participant = _participant_from_cells(row)
        if participant is not None:
            participants.append(participant)
        if len(participants) >= MAX_IMPORT_ROWS:
            break
    return participants


def _parse_xlsx(data: bytes) -> list[ImportParticipant]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl has several format-specific exception classes
        raise ImportParseError("Excel file could not be read") from exc

    participants: list[ImportParticipant] = []
    try:
        sheet = workbook.active
        for row in sheet.iter_rows(values_only=True):
            participant = _participant_from_cells(list(row))
            if participant is not None:
                participants.append(participant)
            if len(participants) >= MAX_IMPORT_ROWS:
                break
    finally:
        workbook.close()
    return participants


def _run_tesseract(image_path: Path) -> str:
    command = [
        "tesseract",
        str(image_path),
        "stdout",
        "-l",
        "Latin+eng+deu+ell+bul",
        "--psm",
        "6",
        "quiet",
    ]
    try:
        result = subprocess.run(command, capture_output=True, text=True, timeout=45, check=False)
    except (FileNotFoundError, subprocess.TimeoutExpired) as exc:
        raise ImportParseError("OCR is not available") from exc
    if result.returncode != 0:
        # Some Tesseract installations do not ship script/Latin. Fall back to core languages.
        fallback = [
            "tesseract", str(image_path), "stdout", "-l", "eng+deu+ell+bul", "--psm", "6", "quiet"
        ]
        result = subprocess.run(fallback, capture_output=True, text=True, timeout=45, check=False)
    if result.returncode != 0:
        raise ImportParseError("Image text could not be recognized")
    return result.stdout


def _ocr_image(data: bytes, suffix: str) -> str:
    with tempfile.TemporaryDirectory(prefix="zahlmeister-import-") as directory:
        image_path = Path(directory) / f"upload{suffix}"
        image_path.write_bytes(data)
        return _run_tesseract(image_path)


def _parse_pdf(data: bytes) -> tuple[list[ImportParticipant], bool]:
    try:
        reader = PdfReader(io.BytesIO(data))
        page_count = min(len(reader.pages), MAX_PDF_PAGES)
        text = "\n".join((reader.pages[index].extract_text() or "") for index in range(page_count))
    except Exception as exc:
        raise ImportParseError("PDF file could not be read") from exc

    participants = _parse_lines(text)
    if participants:
        return participants, False

    with tempfile.TemporaryDirectory(prefix="zahlmeister-pdf-") as directory:
        pdf_path = Path(directory) / "upload.pdf"
        pdf_path.write_bytes(data)
        output_prefix = Path(directory) / "page"
        try:
            result = subprocess.run(
                [
                    "pdftoppm", "-png", "-r", "180", "-f", "1", "-l", str(MAX_PDF_PAGES),
                    str(pdf_path), str(output_prefix),
                ],
                capture_output=True,
                text=True,
                timeout=60,
                check=False,
            )
        except (FileNotFoundError, subprocess.TimeoutExpired) as exc:
            raise ImportParseError("Scanned PDF OCR is not available") from exc
        if result.returncode != 0:
            raise ImportParseError("Scanned PDF could not be rendered")

        ocr_text: list[str] = []
        for image_path in sorted(Path(directory).glob("page-*.png")):
            ocr_text.append(_run_tesseract(image_path))
        return _parse_lines("\n".join(ocr_text)), True


def parse_import(filename: str, content_type: str | None, data: bytes) -> ImportPreview:
    if not data:
        raise ImportParseError("File is empty")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ImportParseError("File is larger than 10 MB")

    suffix = Path(filename or "").suffix.casefold()
    mime = (content_type or "").casefold()
    warnings: list[str] = []

    if suffix in {".xlsx", ".xlsm"}:
        participants = _parse_xlsx(data)
        source_type = "excel"
    elif suffix in {".csv", ".tsv"}:
        participants = _parse_csv(data)
        source_type = "text"
    elif suffix == ".txt" or mime.startswith("text/plain"):
        participants = _parse_lines(_decode_text(data))
        source_type = "text"
    elif suffix == ".pdf" or mime == "application/pdf":
        participants, used_ocr = _parse_pdf(data)
        source_type = "pdf"
        if used_ocr:
            warnings.append("ocr_used")
    elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"} or mime.startswith("image/"):
        participants = _parse_lines(_ocr_image(data, suffix or ".png"))
        source_type = "image"
        warnings.append("ocr_used")
    else:
        raise ImportParseError("File type not supported")

    participants = _deduplicate(participants)
    if not participants:
        raise ImportParseError("No participants could be recognized")
    if len(participants) >= MAX_IMPORT_ROWS:
        warnings.append("row_limit_reached")
    return ImportPreview(source_type=source_type, participants=participants, warnings=warnings)
